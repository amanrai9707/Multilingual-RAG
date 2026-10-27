"""Load the question dataset (from the workbook's Question_Dataset tab, or a CSV export) and CHECK it before freezing.

    python -m src.validate_dataset --dataset HWC_Project_Workbook.xlsx --clean-dir data/clean

Conventions the checker relies on (write them this way in the workbook):
  source_document : exactly the Source_Inventory id, e.g.  WII-001
  source_location : the PDF page(s) of the supporting passage, e.g.  p.12   or   pp.12-13   or   12
                    (same page-numbering convention as src/pdf_to_clean.py - the PDF viewer page number by default)
  terminology_tags: category numbers 1-6 separated by commas, e.g.  2,4      (6 = general language, use it ALONE)
  evidence_grade  : primary | news-2plus | news-1
"""
from __future__ import annotations
import csv, glob, hashlib, os, re
from collections import Counter

FIELDS = ["question_id", "source_document", "source_location", "underlying_fact", "terminology_tags", "english_question",
          "assamese_question", "reference_answer", "supporting_passage", "difficulty", "notes", "evidence_grade"]
VALID_DIFF = {"easy", "medium", "hard"}
VALID_GRADE = {"primary", "news-2plus", "news-1"}
CATEGORIES = {"1": "Species", "2": "Conflict", "3": "Response", "4": "Administrative/SOP", "5": "Geographic/protected area",
              "6": "General language"}


def load_dataset(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True)["Question_Dataset"]
        header = [("" if c.value is None else str(c.value).strip()) for c in ws[1]]
        raw = ws.iter_rows(min_row=2, values_only=True)
        rows = [{h: ("" if v is None else str(v).strip()) for h, v in zip(header, r) if h} for r in raw]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = [{k: (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)]
    return [r for r in rows if r.get("question_id") and not r["question_id"].upper().startswith("EXAMPLE")]


def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_tags(row):
    return set(re.findall(r"\d+", str(row.get("terminology_tags", ""))))


def doc_id_of(row):
    m = re.match(r"\s*([A-Za-z0-9_\-]+)", str(row.get("source_document", "")))
    return m.group(1) if m else ""


def parse_pages(loc):
    """'p.12' -> ['12'];  'pp. 12-14' -> ['12','13','14'];  '12' -> ['12'];  'Section 3, p.9' -> ['9']."""
    s, pages = str(loc), []
    for m in re.finditer(r"(?:pp?\.?|pages?)\s*(\d+)(?:\s*[-\u2013]\s*(\d+))?", s, re.I):
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        pages += [str(x) for x in range(a, min(b, a + 20) + 1)]
    if not pages:
        m = re.fullmatch(r"\s*(\d+)(?:\s*[-\u2013]\s*(\d+))?\s*", s)
        if m:
            a, b = int(m.group(1)), int(m.group(2) or m.group(1))
            pages = [str(x) for x in range(a, min(b, a + 20) + 1)]
    return pages


def gold_keys(row):
    """['WII-001|12', ...] - what retrieval is scored against."""
    d = doc_id_of(row)
    return ["%s|%s" % (d, p) for p in parse_pages(row.get("source_location", ""))]


def _norm(t):
    return re.sub(r"[^a-z0-9]+", "", str(t).lower())


def _assamese_fraction(text):
    letters = [c for c in text if c.isalpha()]
    return (sum("\u0980" <= c <= "\u09ff" for c in letters) / len(letters)) if letters else 0.0


def load_clean_pages(clean_dir):
    from src.ingest import load_clean_doc
    docs = {}
    for path in glob.glob(os.path.join(clean_dir, "*.txt")):
        docs[os.path.splitext(os.path.basename(path))[0]] = {pg: _norm(tx) for pg, tx in load_clean_doc(path)}
    return docs


def validate(rows, clean_dir=None, min_per_category=25, require_assamese=True):
    errors, warnings = [], []
    docs = load_clean_pages(clean_dir) if clean_dir else None
    ids = Counter(r["question_id"] for r in rows)
    for qid, n in ids.items():
        if n > 1:
            errors.append("%s: question_id used %d times" % (qid, n))
    eng = Counter(r.get("english_question", "").strip().lower() for r in rows)
    for q, n in eng.items():
        if q and n > 1:
            warnings.append("duplicate English question (%d times): %s" % (n, q[:70]))
    need = ["source_document", "source_location", "english_question", "reference_answer", "supporting_passage",
            "terminology_tags", "difficulty", "evidence_grade"] + (["assamese_question"] if require_assamese else [])
    for r in rows:
        q = r["question_id"]
        for f in need:
            if not r.get(f):
                errors.append("%s: missing %s" % (q, f))
        tags = parse_tags(r)
        if tags - set(CATEGORIES):
            errors.append("%s: invalid tag(s) %s (use 1-6)" % (q, sorted(tags - set(CATEGORIES))))
        if "6" in tags and len(tags) > 1:
            warnings.append("%s: tag 6 (general language) should be used alone, found %s" % (q, sorted(tags)))
        if r.get("difficulty") and r["difficulty"] not in VALID_DIFF:
            errors.append("%s: difficulty must be easy/medium/hard, got %r" % (q, r["difficulty"]))
        if r.get("evidence_grade") and r["evidence_grade"] not in VALID_GRADE:
            errors.append("%s: evidence_grade must be primary/news-2plus/news-1, got %r" % (q, r["evidence_grade"]))
        if r.get("evidence_grade") == "news-1":
            warnings.append("%s: news-1 (single outlet) - the agreed rule says avoid" % q)
        a, e = r.get("assamese_question", ""), r.get("english_question", "")
        if a:
            if a.strip().lower() == e.strip().lower():
                errors.append("%s: Assamese question is identical to the English one" % q)
            elif _assamese_fraction(a) < 0.5:
                warnings.append("%s: Assamese question does not look like Assamese script (%.0f%% Assamese letters)" % (q, 100 * _assamese_fraction(a)))
        if r.get("evidence_grade") == "primary" and r.get("source_location") and not parse_pages(r["source_location"]):
            errors.append("%s: cannot read a page number from source_location %r (write p.12 or pp.12-13)" % (q, r["source_location"]))
        # checks against the cleaned documents - only for facts claimed as primary
        if docs is not None and r.get("evidence_grade") == "primary":
            d = doc_id_of(r)
            if d not in docs:
                errors.append("%s: no cleaned text file %s.txt in %s" % (q, d, clean_dir))
                continue
            pages = parse_pages(r.get("source_location", ""))
            missing = [p for p in pages if p not in docs[d]]
            if missing:
                errors.append("%s: page(s) %s not present in %s (cleaned doc has %d pages)" % (q, missing, d, len(docs[d])))
            probe = _norm(r.get("supporting_passage", ""))[:80]
            if len(probe) < 20:
                warnings.append("%s: supporting_passage too short to verify" % q)
            elif pages and not any(probe in docs[d].get(p, "") for p in pages):
                elsewhere = [p for p, t in docs[d].items() if probe in t]
                warnings.append("%s: passage not found on cited page(s) %s%s" % (
                    q, pages, " - but it IS on page(s) %s" % elsewhere if elsewhere else " (and not elsewhere in the document: paraphrased, OCR error, or a table?)"))
    cat = Counter()
    for r in rows:
        for t in parse_tags(r):
            cat[t] += 1
    for t in sorted(CATEGORIES):
        if cat[t] < min_per_category:
            warnings.append("category %s (%s): %d questions, floor is %d" % (t, CATEGORIES[t], cat[t], min_per_category))
    counts = {"n_questions": len(rows), "per_category": {t: cat[t] for t in sorted(CATEGORIES)},
              "per_grade": dict(Counter(r.get("evidence_grade", "") for r in rows)),
              "per_difficulty": dict(Counter(r.get("difficulty", "") for r in rows))}
    if not 150 <= len(rows) <= 200:
        warnings.append("%d questions; the protocol target is 150-200" % len(rows))
    return {"errors": errors, "warnings": warnings, "counts": counts}


def format_report(res):
    out = ["Questions: %d | per category: %s | grades: %s | difficulty: %s" % (
        res["counts"]["n_questions"], res["counts"]["per_category"], res["counts"]["per_grade"], res["counts"]["per_difficulty"])]
    out += ["", "ERRORS (%d) - fix before freezing:" % len(res["errors"])] + ["  - " + e for e in res["errors"]]
    out += ["", "WARNINGS (%d) - read, then decide:" % len(res["warnings"])] + ["  - " + w for w in res["warnings"]]
    out += ["", "RESULT: " + ("READY to freeze (no errors)" if not res["errors"] else "NOT ready - %d error(s)" % len(res["errors"]))]
    return "\n".join(out)
